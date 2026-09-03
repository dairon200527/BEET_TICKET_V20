"""GET /api/reportes/rendimiento-convenios/exportar and
GET /api/reportes/afiliados/exportar — real .xlsx downloads (not the
stubbed "not connected to the backend yet" placeholder the frontend used
to show), scoped to the requesting admin's own cooperativa like every
other reportes endpoint."""
from __future__ import annotations

import io

import openpyxl

from tests.conftest import admin_token, auth_headers


def test_exportar_rendimiento_convenios_is_a_real_xlsx(client, admin_a, convenio_a):
    token = admin_token(admin_a)
    r = client.get("/api/reportes/rendimiento-convenios/exportar", headers=auth_headers(token))
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in r.headers["content-disposition"]
    assert "rendimiento_convenios" in r.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "Convenio"

    nombres = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert convenio_a.nombre in nombres


def test_exportar_afiliados_is_a_real_xlsx(client, admin_a, afiliado_a):
    token = admin_token(admin_a)
    r = client.get("/api/reportes/afiliados/exportar", headers=auth_headers(token))
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    documentos = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert afiliado_a.documento in documentos


def test_exportar_is_scoped_to_own_cooperativa(client, admin_b, convenio_a):
    token = admin_token(admin_b)
    r = client.get("/api/reportes/rendimiento-convenios/exportar", headers=auth_headers(token))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    nombres = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert convenio_a.nombre not in nombres


# --- SUPER_ADMIN must be able to export too, same as ADMIN, once a
# cooperativa is selected — regression coverage for the reported bug
# ("las exportaciones del panel de SUPER_ADMIN no funcionan en
# absoluto"). See the final summary for what the real root cause turned
# out to be (these two endpoints were already correctly wired). ---------
def test_superadmin_exportar_rendimiento_sin_cooperativa_es_400(client, superadmin_a):
    token = admin_token(superadmin_a)
    r = client.get("/api/reportes/rendimiento-convenios/exportar", headers=auth_headers(token))
    assert r.status_code == 400


def test_superadmin_exportar_rendimiento_con_cooperativa_seleccionada(client, superadmin_a, convenio_a, cooperativa_a):
    token = admin_token(superadmin_a)
    r = client.get(
        f"/api/reportes/rendimiento-convenios/exportar?cooperativa_id={cooperativa_a.id}", headers=auth_headers(token)
    )
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    nombres = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert convenio_a.nombre in nombres


def test_superadmin_exportar_afiliados_con_cooperativa_seleccionada(client, superadmin_a, afiliado_a, cooperativa_a):
    token = admin_token(superadmin_a)
    r = client.get(f"/api/reportes/afiliados/exportar?cooperativa_id={cooperativa_a.id}", headers=auth_headers(token))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    documentos = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert afiliado_a.documento in documentos
