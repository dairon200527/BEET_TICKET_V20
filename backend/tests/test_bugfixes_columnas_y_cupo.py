"""Acceptance tests for the bug-fix pass:

1. Afiliados carga-masiva creates/updates cupos_credito correctly when
   using human-readable Excel headers (accents, mixed case, spaces
   instead of underscores) — not just the exact snake_case names.
2. A required column still missing after normalization is a clear,
   explicit rejection (400), not a silent skip.
3. A SUPER_ADMIN created via POST /api/admin/usuarios (cooperativa_id=None,
   the normal shape for that role) can actually log in afterward —
   regression test for the AdminOut.cooperativa_id typing bug.

See ../app/utils/bulk_upload.py (normalize_header/require_columns) and
../app/schemas/auth.py (AdminOut) for the actual fixes."""
from __future__ import annotations

from sqlalchemy import select

from app.models.afiliado import Afiliado
from app.models.cupo_credito import CupoCredito
from tests.conftest import admin_token, auth_headers


def _xlsx_upload(headers: list[str], row: list) -> dict:
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return {"file": ("carga.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_carga_masiva_with_human_readable_headers_creates_cupo(client, db_session, admin_a):
    """The exact real-world header style from the bug report: capitalized,
    accented, space-separated — must be read correctly, not silently
    defaulted to zero."""
    token = admin_token(admin_a)
    files = _xlsx_upload(
        ["Documento", "Nombres", "Apellidos", "Correo", "Teléfono", "Estado", "Cupo total", "Cupo disponible"],
        ["4001", "Ana", "Ejemplo", "ana.ejemplo@test.com", "3001111111", "Activo", 500000, 500000],
    )
    r = client.post("/api/afiliados/carga-masiva", files=files, headers=auth_headers(token))
    assert r.status_code == 200
    assert r.json()["creados"] == 1

    afiliado = db_session.execute(select(Afiliado).where(Afiliado.documento == "4001")).scalar_one()
    assert afiliado.telefono == "3001111111"  # accented "Teléfono" was read
    assert afiliado.estado is True

    cupo = db_session.execute(select(CupoCredito).where(CupoCredito.afiliado_id == afiliado.id)).scalar_one()
    assert cupo.cupo_total == 500000
    assert cupo.cupo_disponible == 500000


def test_carga_masiva_cupo_disponible_distinct_from_total_is_honored(client, db_session, admin_a):
    token = admin_token(admin_a)
    files = _xlsx_upload(
        ["Documento", "Nombres", "Apellidos", "Correo", "Cupo total", "Cupo disponible"],
        ["4002", "Con", "Disponible", "con.disponible@test.com", 300000, 100000],
    )
    r = client.post("/api/afiliados/carga-masiva", files=files, headers=auth_headers(token))
    assert r.status_code == 200
    afiliado = db_session.execute(select(Afiliado).where(Afiliado.documento == "4002")).scalar_one()
    cupo = db_session.execute(select(CupoCredito).where(CupoCredito.afiliado_id == afiliado.id)).scalar_one()
    assert cupo.cupo_total == 300000
    assert cupo.cupo_disponible == 100000


def test_carga_masiva_cupo_disponible_blank_defaults_to_cupo_total(client, db_session, admin_a):
    token = admin_token(admin_a)
    files = _xlsx_upload(
        ["Documento", "Nombres", "Apellidos", "Correo", "Cupo total"],
        ["4003", "Sin", "Disponible", "sin.disponible@test.com", 250000],
    )
    r = client.post("/api/afiliados/carga-masiva", files=files, headers=auth_headers(token))
    assert r.status_code == 200
    afiliado = db_session.execute(select(Afiliado).where(Afiliado.documento == "4003")).scalar_one()
    cupo = db_session.execute(select(CupoCredito).where(CupoCredito.afiliado_id == afiliado.id)).scalar_one()
    assert cupo.cupo_total == cupo.cupo_disponible == 250000


def test_carga_masiva_cupo_disponible_mayor_a_total_is_rejected_per_row(client, admin_a):
    token = admin_token(admin_a)
    files = _xlsx_upload(
        ["Documento", "Nombres", "Apellidos", "Correo", "Cupo total", "Cupo disponible"],
        ["4004", "Malo", "Row", "malo@test.com", 100000, 999999],
    )
    r = client.post("/api/afiliados/carga-masiva", files=files, headers=auth_headers(token))
    assert r.status_code == 200
    body = r.json()
    assert body["invalidos"] == 1
    assert "cupo_disponible no puede ser mayor a cupo_total" in body["errores"][0]


def test_carga_masiva_missing_required_column_is_a_clear_400(client, admin_a):
    token = admin_token(admin_a)
    files = _xlsx_upload(["Documento", "Nombres", "Apellidos", "Correo"], ["4005", "X", "Y", "x@test.com"])
    r = client.post("/api/afiliados/carga-masiva", files=files, headers=auth_headers(token))
    assert r.status_code == 400
    assert "cupo_total" in r.json()["detail"]


def test_carga_masiva_duplicate_documento_in_same_file_does_not_crash_the_batch(client, db_session, admin_a):
    """Regression test: a duplicate `documento` within the same uploaded
    file used to crash the WHOLE batch with an IntegrityError on
    cupos_credito's UNIQUE(afiliado_id) — row N's newly-added CupoCredito
    was never flushed, so row N+1's "does a cupo already exist for this
    afiliado" check (on this autoflush=False session) never saw it and
    tried to insert a second one for the same afiliado_id. It resolves
    as a normal upsert (create then update) instead — see
    routers/afiliados.py."""
    token = admin_token(admin_a)
    # Two rows, same documento, distinct cupo_total.
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Documento", "Nombres", "Apellidos", "Correo", "Cupo total"])
    ws.append(["4010", "Primera", "Vez", "primera.vez@test.com", 400000])
    ws.append(["4010", "Primera", "Vez", "primera.vez.v2@test.com", 450000])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = client.post(
        "/api/afiliados/carga-masiva",
        files={"file": ("dup.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers(token),
    )
    assert r.status_code == 200
    body = r.json()
    assert body["creados"] == 1
    assert body["actualizados"] == 1
    assert body["invalidos"] == 0

    afiliado = db_session.execute(select(Afiliado).where(Afiliado.documento == "4010")).scalar_one()
    assert afiliado.correo == "primera.vez.v2@test.com"  # row 2's data won
    cupos = db_session.execute(select(CupoCredito).where(CupoCredito.afiliado_id == afiliado.id)).scalars().all()
    assert len(cupos) == 1  # exactly one cupo row, not a duplicate-key crash


def test_carga_masiva_updates_existing_afiliado_with_human_readable_headers(client, db_session, admin_a, afiliado_a):
    """Same header-normalization robustness, but on the UPDATE path."""
    token = admin_token(admin_a)
    files = _xlsx_upload(
        ["Documento", "Nombres", "Apellidos", "Correo", "Cupo total"],
        [afiliado_a.documento, "Nombre Actualizado", afiliado_a.apellidos, afiliado_a.correo, 600000],
    )
    r = client.post("/api/afiliados/carga-masiva", files=files, headers=auth_headers(token))
    assert r.status_code == 200
    assert r.json()["actualizados"] == 1
    db_session.refresh(afiliado_a)
    assert afiliado_a.nombres == "Nombre Actualizado"
    cupo = db_session.execute(select(CupoCredito).where(CupoCredito.afiliado_id == afiliado_a.id)).scalar_one()
    assert cupo.cupo_total == 600000


# --- convenios/inventario carga-masiva: same normalization, spot-checked ---
def test_convenios_carga_masiva_rejects_missing_required_column(client, admin_a):
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nombre", "Precio publico"])  # missing precio_beet, fecha_inicio
    ws.append(["Convenio X", 10000])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    token = admin_token(admin_a)
    r = client.post(
        "/api/convenios/carga-masiva",
        files={"file": ("c.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers(token),
    )
    assert r.status_code == 400
    assert "precio_beet" in r.json()["detail"]


def test_inventario_carga_masiva_rejects_missing_required_column(client, admin_a):
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Convenio"])  # missing codigo
    ws.append(["Algún convenio"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    token = admin_token(admin_a)
    r = client.post(
        "/api/inventario/carga-masiva",
        files={"file": ("i.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers(token),
    )
    assert r.status_code == 400
    assert "codigo" in r.json()["detail"]


# --- Bug #3: SUPER_ADMIN created via the panel must be able to log in ---
def test_super_admin_created_via_panel_can_log_in(client, db_session, superadmin_sin_cooperativa):
    """Regression test for AdminOut.cooperativa_id being (incorrectly)
    non-nullable — it made login/`/me` raise an unhandled error for
    EXACTLY this case: a SUPER_ADMIN with no cooperativa_id, which is the
    normal, documented shape for that role."""
    token = admin_token(superadmin_sin_cooperativa)

    r_crear = client.post(
        "/api/admin/usuarios",
        json={"nombre": "Nuevo Super Admin", "correo": "nuevo.super.panel@test.com", "password": "ClaveSegura123!", "rol": "SUPER_ADMIN"},
        headers=auth_headers(token),
    )
    assert r_crear.status_code == 201
    assert r_crear.json()["cooperativa_id"] is None

    r_login = client.post(
        "/api/auth/admin/login", json={"correo": "nuevo.super.panel@test.com", "password": "ClaveSegura123!"}
    )
    assert r_login.status_code == 200
    assert r_login.json()["usuario"]["cooperativa_id"] is None

    r_me = client.get("/api/auth/admin/me", headers=auth_headers(r_login.json()["access_token"]))
    assert r_me.status_code == 200
