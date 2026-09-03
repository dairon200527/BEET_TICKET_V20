"""GET /api/convenios/exportar — previously a dead 'Exportar' button in
ConveniosList.jsx with no backend endpoint behind it at all (only the
reportes rendimiento-convenios/afiliados exports existed). Same
resolve_cooperativa_scope pattern as every other export."""
from __future__ import annotations

import io

import openpyxl

from tests.conftest import admin_token, auth_headers


def test_exportar_convenios_is_a_real_xlsx(client, admin_a, convenio_a):
    token = admin_token(admin_a)
    r = client.get("/api/convenios/exportar", headers=auth_headers(token))
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in r.headers["content-disposition"]
    assert "convenios" in r.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "Nombre"

    nombres = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert convenio_a.nombre in nombres


def test_exportar_convenios_scoped_to_own_cooperativa(client, admin_b, convenio_a):
    token = admin_token(admin_b)
    r = client.get("/api/convenios/exportar", headers=auth_headers(token))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    nombres = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert convenio_a.nombre not in nombres


def test_exportar_convenios_superadmin_sin_cooperativa_es_400(client, superadmin_a):
    token = admin_token(superadmin_a)
    r = client.get("/api/convenios/exportar", headers=auth_headers(token))
    assert r.status_code == 400


def test_exportar_convenios_superadmin_con_cooperativa_seleccionada(client, superadmin_a, convenio_a, cooperativa_a):
    token = admin_token(superadmin_a)
    r = client.get(f"/api/convenios/exportar?cooperativa_id={cooperativa_a.id}", headers=auth_headers(token))
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    nombres = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert convenio_a.nombre in nombres
