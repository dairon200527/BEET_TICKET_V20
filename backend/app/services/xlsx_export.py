"""
Shared .xlsx export builder for report download endpoints
(GET /api/reportes/.../exportar, GET /api/afiliados/exportar).

Every export goes through `build_xlsx` so the visual style (header
color, font, column widths) never drifts between reports — the same
reasoning as `bulk_upload.parse_rows` being the one shared reader for
every bulk-upload endpoint.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_xlsx(sheet_name: str, headers: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]  # Excel's own sheet-name length limit

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1C26E5", end_color="1C26E5", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    body_font = Font(name="Arial", size=11)
    for row in rows:
        ws.append(row)
    for row_idx in range(2, len(rows) + 2):
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).font = body_font

    for col_idx, header in enumerate(headers, start=1):
        width = max(12, min(40, len(header) + 4))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def dated_filename(base: str) -> str:
    return f"{base}_{date.today().isoformat()}.xlsx"


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
